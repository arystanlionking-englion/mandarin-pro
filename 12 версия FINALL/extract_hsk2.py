import openpyxl
import json

wb = openpyxl.load_workbook(r'D:\qwerty\8 версия\2 hsk диалоги и слова\СЛОВА\HSK 2 Слова.xlsx', data_only=True)
ws = wb['HSK 1']

in_hsk2 = False
current_lesson = None
words = []
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
    vals = [str(v) if v is not None else '' for v in row]
    if 'HSK 2' in vals[0] or 'HSK 2' in ' '.join(vals):
        in_hsk2 = True
        continue
    if not in_hsk2:
        continue
    if vals[0].startswith('Lesson '):
        try:
            current_lesson = int(vals[0].replace('Lesson ', ''))
        except:
            continue
        continue
    if vals[0].isdigit() and current_lesson is not None:
        num = int(vals[0])
        hanzi = vals[1]
        pinyin = vals[2]
        pos = vals[3]
        en = vals[4]
        ru = vals[5]
        kz = vals[6]
        if hanzi:
            words.append({
                'lesson': current_lesson,
                'hanzi': hanzi,
                'pinyin': pinyin,
                'pos': pos,
                'en': en,
                'ru': ru,
                'kz': kz
            })

print(f'Total HSK2 words: {len(words)}')
print(f'Lessons: {sorted(set(w["lesson"] for w in words))}')
print()

# Print as JS array
for w in words:
    # Escape single quotes
    for k in w:
        if isinstance(w[k], str):
            w[k] = w[k].replace("'", "\\'")
    print(f"  {{lesson:{w['lesson']},hanzi:'{w['hanzi']}',pinyin:'{w['pinyin']}',pos:'{w['pos']}',en:'{w['en']}',ru:'{w['ru']}',kz:'{w['kz']}'}},")