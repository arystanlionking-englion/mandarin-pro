#!/usr/bin/env python3
"""Extract HSK 4A data from xlsx/docx and generate JavaScript constants. V2"""
import json
import openpyxl
from docx import Document
import os
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

BASE = r"D:\qwerty\11 версия FINALL"
XLSX_PATH = os.path.join(BASE, "4 hsk 1 A -  1 часть диалоги и слова", "Слова", "HSK 4Слова.xlsx")
DIALOG_DIR = os.path.join(BASE, "4 hsk 1 A -  1 часть диалоги и слова", "Диалоги")
OUTPUT_PATH = os.path.join(BASE, "hsk4a_data.js")

# ===== 1. Extract words from xlsx =====
print("=== Extracting words ===")
wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws = wb[wb.sheetnames[0]]

words = []
current_lesson = 0

for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
    a = row[0] if len(row) > 0 else None
    b = row[1] if len(row) > 1 else None
    c = row[2] if len(row) > 2 else None
    d = row[3] if len(row) > 3 else None
    e = row[4] if len(row) > 4 else None
    f = row[5] if len(row) > 5 else None
    g = row[6] if len(row) > 6 else None

    if a and isinstance(a, str) and a.startswith("Lesson"):
        try:
            current_lesson = int(a.replace("Lesson", "").strip())
        except:
            pass
        continue

    if a == '№' or a == 'HSK 4' or a is None:
        continue

    if b and c and isinstance(b, str) and isinstance(c, str):
        words.append({
            "lesson": current_lesson,
            "hanzi": str(b).strip(),
            "pinyin": str(c).strip(),
            "pos": str(d).strip() if d else "",
            "en": str(e).strip() if e else "",
            "ru": str(f).strip() if f else "",
            "kz": str(g).strip() if g else ""
        })

wb.close()
print(f"Words: {len(words)}, Lessons: {sorted(set(w['lesson'] for w in words))}")

# ===== 2. Extract dialogues from docx =====
print("\n=== Extracting dialogues ===")

def strip_bold(text):
    """Remove ** markers from text."""
    return text.replace('**', '').strip()

def extract_speaker_text(line):
    """Extract speaker and text from a dialogue line. Returns (speaker, text) or None."""
    line = line.strip()
    if not line:
        return None
    
    # Pattern: * **speaker：** text  (bullet + bold speaker)
    m = re.match(r'^\*\s+\*\*([^*]+?)[：:]\*\*\s*(.+)', line)
    if m:
        return (m.group(1).strip(), m.group(2).strip())
    
    # Pattern: **speaker：** text  (bold speaker, no bullet)
    m = re.match(r'^\*\*([^*]+?)[：:]\*\*\s*(.+)', line)
    if m:
        return (m.group(1).strip(), m.group(2).strip())
    
    # Pattern: speaker：text  (plain, for simpler formats)
    m = re.match(r'^([A-Za-z\u4e00-\u9fff\u3400-\u4dbf]{1,10})[：:]\s*(.+)', line)
    if m:
        speaker = m.group(1).strip()
        text = m.group(2).strip()
        if len(speaker) <= 6 and re.search(r'[\u4e00-\u9fff]', text):
            return (speaker, text)
    
    return None

def is_dialogue_header(text):
    """Check if text is a dialogue/section header. Returns cleaned title or None."""
    text = text.strip()
    if not text:
        return None
    
    # Remove leading/trailing ** markers
    clean = text.replace('**', '').strip()
    
    # Pattern: **N. Title (description)**
    m = re.match(r'^(\d+)\.\s*(.+)', clean)
    if m:
        title = m.group(2).strip()
        # Extract Chinese part
        cn = re.search(r'([\u4e00-\u9fff][\u4e00-\u9fff\s\w·。！？，、：；""''（）]*)', title)
        if cn:
            return cn.group(1).strip()[:40]
        return title[:40]
    
    # Pattern: **Диалог N: Title (description)**
    m = re.match(r'^Диалог\s*\d+[：:]\s*(.+)', clean)
    if m:
        title = m.group(1).strip()
        cn = re.search(r'([\u4e00-\u9fff][\u4e00-\u9fff\s\w·。！？，、：；""''（）]*)', title)
        if cn:
            return cn.group(1).strip()[:40]
        return title[:40]
    
    # Pattern: **课文 N: Title**
    m = re.match(r'^课文\s*\d+[：:]\s*(.+)', clean)
    if m:
        title = m.group(1).strip()
        cn = re.search(r'([\u4e00-\u9fff][\u4e00-\u9fff\s\w·。！？，、：；""''（）]*)', title)
        if cn:
            return cn.group(1).strip()[:40]
        return title[:40]
    
    # Pattern: N-M Title (like "05-1 王静在家具店买沙发")
    m = re.match(r'^\d{2}-\d+\s+(.+)', clean)
    if m:
        return m.group(1).strip()[:40]
    
    # Pattern: **Диалоги (Dialogues)** or **Тексты (Texts)**
    if re.match(r'^(Диалоги|Тексты|Часть)', clean):
        return None  # Section header, not a dialogue
    
    return None

def is_section_marker(text):
    """Check if text marks a section boundary (Тексты, etc.)."""
    clean = text.replace('**', '').strip()
    return bool(re.match(r'^(Тексты|Часть\s*\d+|##\s*Тексты)', clean))

def is_text_content(text):
    """Check if text is a reading text (not dialogue)."""
    clean = text.replace('**', '').strip()
    # Pattern: **课文 N** followed by long text
    if re.match(r'^课文\s*\d+\s*$', clean):
        return True
    return False

def extract_dialogues(filepath, lesson_num):
    """Extract all dialogues from a docx file."""
    doc = Document(filepath)
    all_text = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            all_text.append(t)
    
    dialogues = []
    current_title = ""
    current_lines = []
    in_text_section = False
    
    for text in all_text:
        # Skip intro/metadata lines
        if any(text.startswith(p) for p in [
            "Вот ", "Ниже ", "Конечно", "Понял тебя", "Вот текст",
            "Вот оригинальный", "Ниже представлен"
        ]):
            continue
        if "image.png" in text or "image_" in text:
            continue
        
        # Check for text section marker (reading texts, not dialogues)
        if is_section_marker(text):
            # Save any current dialogue
            if current_lines and not in_text_section:
                dialogues.append({
                    "type": "dialogue",
                    "title": current_title or f"Диалог {len(dialogues)+1}",
                    "lines": current_lines
                })
                current_lines = []
            in_text_section = True
            continue
        
        # Skip ## header lines that are just section markers
        if text.startswith('## ') and not is_dialogue_header(text):
            clean = text.replace('## ', '').strip()
            if re.match(r'^(Диалоги|Тексты|Часть|💬)', clean):
                continue
        
        # Check for dialogue header
        title = is_dialogue_header(text)
        if title is not None:
            # Save previous dialogue
            if current_lines:
                dialogues.append({
                    "type": "dialogue",
                    "title": current_title or f"Диалог {len(dialogues)+1}",
                    "lines": current_lines
                })
            current_title = title
            current_lines = []
            in_text_section = False
            continue
        
        # If in text section, skip until next dialogue
        if in_text_section:
            continue
        
        # Try to extract speaker line
        sp = extract_speaker_text(text)
        if sp:
            current_lines.append({"speaker": sp[0], "text": sp[1]})
            continue
        
        # Stage direction or continuation
        if current_lines and re.search(r'[\u4e00-\u9fff]', text) and not text.startswith('#'):
            # Parenthetical = stage direction
            if text.startswith('(') or text.startswith('（'):
                current_lines.append({"speaker": "", "text": text})
            elif len(text) < 200:
                current_lines.append({"speaker": "", "text": text})
    
    # Save last dialogue
    if current_lines and not in_text_section:
        dialogues.append({
            "type": "dialogue",
            "title": current_title or f"Диалог {len(dialogues)+1}",
            "lines": current_lines
        })
    
    # Filter: keep only dialogues with 2+ lines that have speakers
    result = []
    for d in dialogues:
        # Count lines with actual speakers
        speaker_lines = [l for l in d['lines'] if l['speaker']]
        if len(speaker_lines) >= 2:
            result.append(d)
    
    return result

# Process all docx files
dialogues_by_lesson = {}
dialogue_files = sorted(
    [f for f in os.listdir(DIALOG_DIR) if f.endswith('.docx')],
    key=lambda x: int(re.match(r'(\d+)', x).group(1)) if re.match(r'(\d+)', x) else 0
)

for fname in dialogue_files:
    lesson_match = re.match(r'(\d+)', fname)
    if not lesson_match:
        continue
    lesson_num = int(lesson_match.group(1))
    
    filepath = os.path.join(DIALOG_DIR, fname)
    dials = extract_dialogues(filepath, lesson_num)
    dialogues_by_lesson[str(lesson_num)] = dials
    
    print(f"\nLesson {lesson_num}: {len(dials)} dialogues")
    for d in dials:
        print(f"  [{len(d['lines'])} lines] {d['title']}")

# ===== 3. Generate Lessons metadata =====
lessons = []
lesson_data = {
    1: ("爱情与生活", "Любовь и жизнь", "Махаббат пен өмір"),
    2: ("朋友与适应", "Друзья и адаптация", "Достар мен бейімделу"),
    3: ("面试与工作", "Собеседование и работа", "Сұхбат пен жұмыс"),
    4: ("工作与收入", "Работа и доход", "Жұмыс мен табыс"),
    5: ("购物与家具", "Покупки и мебель", "Сатып алу мен жиһаз"),
    6: ("买衣服", "Покупка одежды", "Киім сатып алу"),
    7: ("天气与健康", "Погода и здоровье", "Ауа райы мен денсаулық"),
    8: ("美食与运动", "Еда и спорт", "Тамақ мен спорт"),
    9: ("减肥与坚持", "Похудение и настойчивость", "Арықтау мен табандылық"),
    10: ("家庭与幸福", "Семья и счастье", "Отбасы мен бақыт"),
    11: ("学习与考试", "Учёба и экзамены", "Оқу мен емтихан"),
    12: ("生意与经验", "Бизнес и опыт", "Бизнес мен тәжірибе"),
    13: ("京剧与艺术", "Пекинская опера и искусство", "Пекин операсы мен өнер"),
    14: ("出差与环保", "Командировка и экология", "Іссапар мен экология"),
    15: ("教育与习惯", "Воспитание и привычки", "Тәрбие мен әдет"),
    16: ("留学与签证", "Учёба за рубежом и виза", "Шетелде оқу мен виза"),
    17: ("秋天与自然", "Осень и природа", "Күз мен табиғат"),
    18: ("科学与技术", "Наука и технологии", "Ғылым мен технология"),
    19: ("申请与表格", "Заявления и документы", "Өтініш пен құжаттар"),
    20: ("旅行与计划", "Путешествия и планы", "Саяхат пен жоспарлар"),
}

for i in range(1, 21):
    cn, ru, kz = lesson_data.get(i, (f"Урок {i}", f"Урок {i}", f"Сабақ {i}"))
    lessons.append({"id": i, "cn": cn, "ru": ru, "kz": kz})

# ===== 4. Generate phrases from words =====
phrases = []
for w in words:
    phrases.append({
        "lesson": w["lesson"],
        "hanzi": w["hanzi"],
        "pinyin": w["pinyin"],
        "ru": w["ru"],
        "kz": w["kz"]
    })

# ===== 5. Output JavaScript =====
print(f"\n=== Generating {OUTPUT_PATH} ===")

def js_str(s):
    """Escape string for JavaScript single-quoted literal."""
    return s.replace("\\", "\\\\").replace("'", "\\'").replace("\n", "\\n")

with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
    f.write("// ===== HSK 4A WORDS =====\n")
    f.write("const HSK4A_WORDS_RAW = [\n")
    for w in words:
        f.write(f"  {{lesson:{w['lesson']},hanzi:'{js_str(w['hanzi'])}',pinyin:'{js_str(w['pinyin'])}',pos:'{js_str(w['pos'])}',en:'{js_str(w['en'])}',ru:'{js_str(w['ru'])}',kz:'{js_str(w['kz'])}'}},\n")
    f.write("];\n\n")
    
    f.write("// ===== HSK 4A PHRASES =====\n")
    f.write("const HSK4A_PHRASES = [\n")
    for p in phrases:
        f.write(f"  {{lesson:{p['lesson']},hanzi:'{js_str(p['hanzi'])}',pinyin:'{js_str(p['pinyin'])}',ru:'{js_str(p['ru'])}',kz:'{js_str(p['kz'])}'}},\n")
    f.write("];\n\n")
    
    f.write("// ===== HSK 4A LESSONS =====\n")
    f.write("const HSK4A_LESSONS = [\n")
    for l in lessons:
        f.write(f"  {{id:{l['id']},cn:'{js_str(l['cn'])}',ru:'{js_str(l['ru'])}',kz:'{js_str(l['kz'])}'}},\n")
    f.write("];\n\n")
    
    f.write("// ===== HSK 4A DIALOGUES =====\n")
    f.write("const HSK4A_DIALOGUES = {\n")
    for lesson_key in sorted(dialogues_by_lesson.keys(), key=int):
        dials = dialogues_by_lesson[lesson_key]
        if not dials:
            continue
        f.write(f'  {lesson_key}: [\n')
        for d in dials:
            f.write(f"    {{type:'dialogue',title:'{js_str(d['title'])}',lines:[\n")
            for line in d['lines']:
                f.write(f"      {{speaker:'{js_str(line['speaker'])}',text:'{js_str(line['text'])}'}},\n")
            f.write(f"    ]}},\n")
        f.write(f"  ],\n")
    f.write("};\n")

print("\nDone!")
print(f"  Words: {len(words)}")
print(f"  Phrases: {len(phrases)}")
print(f"  Lessons: {len(lessons)}")
dialogue_lessons = {k: v for k, v in dialogues_by_lesson.items() if v}
print(f"  Dialogue lessons: {len(dialogue_lessons)}")
total_dial = sum(len(v) for v in dialogues_by_lesson.values())
print(f"  Total dialogues: {total_dial}")
print(f"  Dialogue coverage: {sorted(int(k) for k in dialogue_lessons.keys())}")